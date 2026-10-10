"""The KPI workbook's shape and its rules, pinned without a database.

`app/kpi_workbook.py` reads a deployment and writes a spreadsheet the office
will act on, so what matters is pinned here on hand-built rows: the LinkedIn
classifier agrees with the register form, the buckets have closed edges, the
streak rule is the student screen's, every feature names a screen that still
exists, the sign-in doors are the router's, the data sheets carry no names
unless asked, and every formula points at a sheet that is in the file. One
test at the end reads the seeded database when it is reachable, which is the
only way to prove the SELECTs against the real schema.
"""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart

from app import kpi_workbook as kw
from app.clock import programme_tz
from tests.conftest import requires_db

REPO = Path(__file__).resolve().parent.parent.parent.parent
ROUTES = REPO / "apps" / "web" / "src" / "app" / "app.routes.ts"

TZ = programme_tz()
TAKEN_AT = datetime(2026, 10, 10, 9, 30, tzinfo=TZ)
TODAY = TAKEN_AT.date()


def _student(i: int, **over) -> kw.StudentFacts:
    base = dict(
        student_id=f"stu{i}",
        user_id=f"usr{i}",
        name=f"Student {i}",
        usn=f"1MP25MDM{i:02d}",
        email=f"1mp25mdm{i:02d}@bgscet.ac.in",
        account_created_at=TAKEN_AT - timedelta(days=120),
        last_login_at=None,
        google_linked=True,
        password_set=False,
        disabled=False,
        stage="EXCEL",
        semester=2,
        seated=True,
        batch="MBA - Finance · 2025-27",
        department="Management Studies",
        course="MBA",
        specialization="Finance",
        mentor_assigned=True,
        profile_on_record=True,
        fields_on_file=frozenset({"Phone", "LinkedIn", "City"}),
        linkedin_url="https://www.linkedin.com/in/student",
        placement_eligible=True,
        interested_in_jobs=True,
        interested_in_internships=False,
        profile_updated_at=TAKEN_AT - timedelta(days=3),
        marks_on_record=True,
        attendance_on_record=False,
        feature_uses={f.key: (0, 0) for f in kw.FEATURES},
        mock_interviews_completed=0,
    )
    base.update(over)
    return kw.StudentFacts(**base)


def _registration(**over) -> kw.RegistrationFacts:
    base = dict(
        name="Asha Rao",
        email="1bg24mba045@bgscet.ac.in",
        created_at=TAKEN_AT - timedelta(days=10),
        status="APPROVED",
        reviewed_at=TAKEN_AT - timedelta(days=9),
        degree_level="PG",
        department="Management Studies",
        course="MBA",
        specialization="Finance",
        dual=False,
        has_usn=True,
        has_phone=True,
        has_personal_email=True,
        has_linkedin=True,
        has_cv=True,
        has_photo=True,
        repeat_application=False,
        became_student=True,
    )
    base.update(over)
    return kw.RegistrationFacts(**base)


def _snapshot(students=None, registrations=None, login_days=None, sign_ins=None, window_days=30) -> kw.Snapshot:
    return kw.Snapshot(
        taken_at=TAKEN_AT,
        window_days=window_days,
        source="test",
        students=list(students or []),
        graduated=2,
        removed=1,
        registrations=list(registrations or []),
        login_days=list(login_days or []),
        sign_ins=list(sign_ins or []),
    )


def _three_students() -> kw.Snapshot:
    """A daily student, an occasional one and one who never signed in."""
    daily_days = [("usr1", TODAY - timedelta(days=d)) for d in range(0, 20)]
    occasional_days = [("usr2", TODAY - timedelta(days=d)) for d in (3, 9, 40, 41)]
    uses = {f.key: (0, 0) for f in kw.FEATURES}
    uses["mock_interview"] = (4, 2)
    uses["ledger"] = (30, 12)
    uses["assistant"] = (7, 1)
    students = [
        _student(1, feature_uses=uses, mock_interviews_completed=3),
        _student(2, linkedin_url="https://www.linkedin.com/company/bgscet", fields_on_file=frozenset({"LinkedIn"}),
                 feature_uses={**{f.key: (0, 0) for f in kw.FEATURES}, "upload": (2, 0)}),
        _student(3, profile_on_record=False, linkedin_url=None, fields_on_file=frozenset(), seated=False,
                 batch=None, course=None, specialization=None, mentor_assigned=False,
                 placement_eligible=None, interested_in_jobs=None, interested_in_internships=None,
                 profile_updated_at=None, marks_on_record=False),
    ]
    sign_ins = [
        kw.SignIn("usr1", datetime.combine(TODAY, datetime.min.time(), tzinfo=TZ).replace(hour=21), "google",
                  "Mozilla/5.0 (Linux; Android 14; SM-A155F) Chrome/128.0 Mobile Safari/537.36"),
        kw.SignIn("usr1", datetime.combine(TODAY - timedelta(days=1), datetime.min.time(), tzinfo=TZ).replace(hour=9),
                  "password", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/128.0 Safari/537.36"),
        kw.SignIn("usr2", datetime.combine(TODAY - timedelta(days=3), datetime.min.time(), tzinfo=TZ).replace(hour=14),
                  "code", None),
    ]
    registrations = [
        _registration(),
        _registration(email="x@bgscet.ac.in", status="REJECTED", became_student=False, repeat_application=True),
        _registration(email="x@bgscet.ac.in", status="PENDING_REVIEW", reviewed_at=None, became_student=False,
                      repeat_application=True, created_at=TAKEN_AT - timedelta(days=2), department=None),
        _registration(email="y@bgscet.ac.in", status="AUTO_APPROVED", reviewed_at=None,
                      created_at=TAKEN_AT - timedelta(days=400), has_linkedin=False, has_cv=False),
    ]
    return _snapshot(students, registrations, daily_days + occasional_days, sign_ins)


# --------------------------------------------------------------------------- #
# The pure rules.
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    "url, expected",
    [
        (None, kw.LINKEDIN_MISSING),
        ("   ", kw.LINKEDIN_MISSING),
        ("https://www.linkedin.com/in/asha-rao", kw.LINKEDIN_PERSONAL),
        ("linkedin.com/in/asha", kw.LINKEDIN_PERSONAL),
        ("https://in.linkedin.com/in/asha?trk=share", kw.LINKEDIN_PERSONAL),
        ("https://m.linkedin.com/in/asha/", kw.LINKEDIN_PERSONAL),
        ("https://www.linkedin.com/company/bgscet", kw.LINKEDIN_OTHER_PAGE),
        ("https://www.linkedin.com/school/bgscet/", kw.LINKEDIN_OTHER_PAGE),
        ("https://www.linkedin.com/posts/asha_activity-1", kw.LINKEDIN_OTHER_PAGE),
        ("https://www.linkedin.com/in/", kw.LINKEDIN_OTHER_PAGE),
        ("https://lnkd.in/gAbc123", kw.LINKEDIN_SHORT),
        ("https://www.linkedin.com/", kw.LINKEDIN_NOT_LINKEDIN),
        ("instagram.com/asha.rao", kw.LINKEDIN_NOT_LINKEDIN),
        ("linkedin", kw.LINKEDIN_NOT_LINKEDIN),
        ("https://notlinkedin.com/in/asha", kw.LINKEDIN_NOT_LINKEDIN),
    ],
)
def test_linkedin_boxes_are_classified_by_what_they_hold(url, expected) -> None:
    assert kw.classify_linkedin(url) == expected


def test_every_link_the_register_form_accepts_is_a_linkedin_link_here() -> None:
    """The form's validator and this classifier share one host rule; a link
    that passes `/register` must never read as 'Not a LinkedIn link' on the
    sheet, or the office is told a student lied on a box the API checked."""
    from app.routers.registration import RegisterIn

    accepted = [
        "linkedin.com/in/asha",
        "www.linkedin.com/in/asha",
        "https://in.linkedin.com/in/asha-rao-1a2b3c",
        "https://m.linkedin.com/in/asha",
        "https://lnkd.in/gXyZ",
        "https://www.linkedin.com/company/bgscet",
    ]
    for raw in accepted:
        stored = RegisterIn._linkedin_shape(raw)
        status = kw.classify_linkedin(stored)
        assert status not in (kw.LINKEDIN_MISSING, kw.LINKEDIN_NOT_LINKEDIN), (raw, stored, status)


@pytest.mark.parametrize(
    "agent, family",
    [
        (None, kw.DEVICE_OTHER),
        ("", kw.DEVICE_OTHER),
        ("Mozilla/5.0 (Linux; Android 14; SM-A155F) AppleWebKit/537.36 Chrome/128.0 Mobile Safari/537.36", kw.DEVICE_ANDROID),
        ("Mozilla/5.0 (iPhone; CPU iPhone OS 17_5 like Mac OS X) Safari/604.1", kw.DEVICE_IOS),
        ("Mozilla/5.0 (iPad; CPU OS 17_5 like Mac OS X) Safari/604.1", kw.DEVICE_IOS),
        ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/128.0 Safari/537.36", kw.DEVICE_WINDOWS),
        ("Mozilla/5.0 (Macintosh; Intel Mac OS X 14_5) Safari/605.1.15", kw.DEVICE_MAC),
        ("Mozilla/5.0 (X11; Linux x86_64) Chrome/128.0 Safari/537.36", kw.DEVICE_OTHER),
        ("curl/8.5.0", kw.DEVICE_OTHER),
    ],
)
def test_device_family_is_coarse_and_never_raises(agent, family) -> None:
    assert kw.device_family(agent) == family


def test_frequency_and_recency_buckets_have_closed_edges() -> None:
    assert [kw.frequency_bucket(n) for n in (0, 1, 2, 3, 7, 8, 15, 16, 30)] == [
        kw.FREQ_NONE, kw.FREQ_1_2, kw.FREQ_1_2, kw.FREQ_3_7, kw.FREQ_3_7,
        kw.FREQ_8_15, kw.FREQ_8_15, kw.FREQ_16_PLUS, kw.FREQ_16_PLUS,
    ]
    assert [kw.recency_bucket(n) for n in (None, 0, 1, 7, 8, 30, 31, 90, 91)] == [
        kw.RECENCY_NEVER, kw.RECENCY_TODAY, kw.RECENCY_WEEK, kw.RECENCY_WEEK, kw.RECENCY_MONTH,
        kw.RECENCY_MONTH, kw.RECENCY_DORMANT, kw.RECENCY_DORMANT, kw.RECENCY_LOST,
    ]
    # Every bucket a pie can draw is in its tuple, so a slice is never unlabelled.
    for n in range(0, 40):
        assert kw.frequency_bucket(n) in kw.FREQUENCY_BUCKETS
        assert kw.recency_bucket(n) in kw.RECENCY_BUCKETS


def test_completeness_bands_follow_the_twelve_fields() -> None:
    total = len(kw.PROFILE_FIELDS)
    assert total == 12
    assert kw.completeness_bucket(False, 0) == kw.COMPLETE_NONE
    assert kw.completeness_bucket(True, 12) == kw.COMPLETE_FULL
    assert kw.completeness_bucket(True, 11) == kw.COMPLETE_FULL  # 91.7%
    assert kw.completeness_bucket(True, 10) == kw.COMPLETE_MOST  # 83.3%
    assert kw.completeness_bucket(True, 4) == kw.COMPLETE_STARTED  # 33.3%
    assert kw.completeness_bucket(True, 3) == kw.COMPLETE_BARELY  # 25%
    assert kw.completeness_bucket(True, 0) == kw.COMPLETE_BARELY


def test_streaks_count_the_way_the_student_screen_does() -> None:
    today = date(2026, 10, 10)
    assert kw.streaks([], today) == (0, 0)
    # Ends today: current is the run up to today.
    days = [today - timedelta(days=d) for d in range(5)] + [today - timedelta(days=20), today - timedelta(days=21)]
    assert kw.streaks(days, today) == (5, 5)
    # Ends yesterday: today is not opened yet and does not break it.
    assert kw.streaks([today - timedelta(days=d) for d in (1, 2, 3)], today) == (3, 3)
    # Ended the day before yesterday: no current streak, the longest survives.
    assert kw.streaks([today - timedelta(days=d) for d in (2, 3, 4, 5)], today) == (0, 4)
    # Duplicates (two sign-ins on one day) are one day.
    assert kw.streaks([today, today, today - timedelta(days=1)], today) == (2, 2)


def test_every_feature_names_a_student_screen_that_still_exists() -> None:
    """A KPI row that names a deleted screen reads as a working feature nobody
    uses. The route table is the truth about which screens exist."""
    routes = ROUTES.read_text(encoding="utf-8")
    for feature in kw.FEATURES:
        assert f"path: '{feature.route}'" in routes, f"{feature.key} names {feature.route}, which is not routed"
    assert len({f.key for f in kw.FEATURES}) == len(kw.FEATURES)
    assert len({f.label for f in kw.FEATURES}) == len(kw.FEATURES)


def test_the_door_vocabulary_is_the_sign_in_routers() -> None:
    from app.routers import auth

    assert set(kw.DOORS) == {auth.DOOR_PASSWORD, auth.DOOR_CODE, auth.DOOR_GOOGLE, auth.DOOR_ACTIVATION}
    assert set(kw.DOOR_LABELS) == set(kw.DOORS)


def test_the_pie_palette_never_needs_a_seventh_colour() -> None:
    """Every fixed category list a pie draws fits the six-slot palette; the
    two data-driven ones (department, feature share) fold into Other at five."""
    for buckets in (kw.FREQUENCY_BUCKETS, kw.RECENCY_BUCKETS, kw.LINKEDIN_STATUSES, kw.COMPLETENESS_BUCKETS,
                    kw.REGISTRATION_STATUS_ORDER, kw.DEVICES, tuple(kw.STAGE_LABELS.values())):
        assert len(buckets) <= len(kw.PALETTE)
    assert len(kw.DOORS) + 1 <= len(kw.PALETTE)  # the doors plus Other


# --------------------------------------------------------------------------- #
# The workbook.
# --------------------------------------------------------------------------- #


@pytest.fixture(scope="module")
def workbook_paths(tmp_path_factory) -> tuple[Path, Path]:
    folder = tmp_path_factory.mktemp("kpis")
    plain = kw.write_workbook(_three_students(), folder / "plain.xlsx")
    named = kw.write_workbook(_three_students(), folder / "named.xlsx", roster=True)
    return plain, named


def test_the_workbook_has_every_promised_sheet_in_reading_order(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    assert wb.sheetnames == [
        kw.SHEET_SUMMARY, kw.SHEET_SIGN_INS, kw.SHEET_REGISTRATION, kw.SHEET_FUNCTIONALITY, kw.SHEET_PROFILE,
        kw.SHEET_PAGE_VIEWS, kw.SHEET_DEFINITIONS, kw.SHEET_DATA_STUDENTS, kw.SHEET_DATA_REGISTRATIONS,
        kw.SHEET_DATA_SIGN_INS,
    ]
    assert wb.calculation.fullCalcOnLoad is True


def test_every_kpi_sheet_draws_pies_of_at_most_six_slices(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    pies = 0
    for name in (kw.SHEET_SIGN_INS, kw.SHEET_REGISTRATION, kw.SHEET_FUNCTIONALITY, kw.SHEET_PROFILE):
        charts = wb[name]._charts
        assert charts, f"{name} draws no chart"
        for chart in charts:
            assert isinstance(chart, (PieChart, BarChart))
            if isinstance(chart, PieChart):
                pies += 1
                ref = chart.series[0].val.numRef.f
                first, last = (int(m) for m in re.findall(r"\$(\d+)", ref))
                assert last - first + 1 <= len(kw.PALETTE), ref
                # Every slice carries its palette colour, in table order.
                assert [pt.idx for pt in chart.series[0].dPt] == list(range(last - first + 1))
    assert pies >= 8  # frequency, recency, door, device, status, department, share of use, LinkedIn, completeness, stage


def test_the_data_sheets_carry_no_names_unless_asked(workbook_paths) -> None:
    plain, named = (load_workbook(p) for p in workbook_paths)
    students_plain = [c.value for c in plain[kw.SHEET_DATA_STUDENTS][1]]
    students_named = [c.value for c in named[kw.SHEET_DATA_STUDENTS][1]]
    for column in kw.ROSTER_STUDENT_COLUMNS:
        assert column not in students_plain
        assert column in students_named
    regs_plain = [c.value for c in plain[kw.SHEET_DATA_REGISTRATIONS][1]]
    regs_named = [c.value for c in named[kw.SHEET_DATA_REGISTRATIONS][1]]
    for column in kw.ROSTER_REGISTRATION_COLUMNS:
        assert column not in regs_plain
        assert column in regs_named
    # And no cell anywhere in the plain file holds a name, USN or address.
    for ws in plain.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert "Student 1" not in value and "bgscet.ac.in" not in value and "1MP25MDM" not in value, value
    assert "no names" in str(plain[kw.SHEET_SUMMARY]["A3"].value)
    assert "--roster" in str(named[kw.SHEET_SUMMARY]["A3"].value)


def test_every_formula_points_at_a_sheet_in_the_file(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    pattern = re.compile(r"'([^']+)'!")
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                    for sheet in pattern.findall(cell.value):
                        assert sheet in wb.sheetnames, (ws.title, cell.coordinate, cell.value)
                    # Nothing LibreOffice cannot evaluate (the xlsx skill's list).
                    assert not re.search(r"\b(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SEQUENCE)\(", cell.value), cell.value
    assert formulas > 100


def test_the_summary_is_formulas_over_the_other_sheets(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    ws = wb[kw.SHEET_SUMMARY]
    values = [(row[1].value, row[2].value) for row in ws.iter_rows(min_row=6) if row[1].value and row[2].value]
    assert values, "the summary table is empty"
    for label, value in values:
        if label == "Most visited pages":
            assert value == "not recorded"
        elif label in {kw.SHEET_SIGN_INS, kw.SHEET_REGISTRATION, kw.SHEET_FUNCTIONALITY, kw.SHEET_PROFILE,
                       kw.SHEET_PAGE_VIEWS, kw.SHEET_DEFINITIONS, kw.SHEET_DATA_STUDENTS,
                       kw.SHEET_DATA_REGISTRATIONS, kw.SHEET_DATA_SIGN_INS}:
            continue  # the sheet index at the foot
        else:
            assert isinstance(value, str) and value.startswith("="), (label, value)


def test_the_students_sheet_says_what_the_rows_say(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    ws = wb[kw.SHEET_DATA_STUDENTS]
    headers = [c.value for c in ws[1]]
    rows = {row[0]: dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {1, 2, 3}
    daily, occasional, never = rows[1], rows[2], rows[3]
    assert daily["Sign-in days (last 30 days)"] == 20
    assert daily["Sign-in frequency"] == kw.FREQ_16_PLUS
    assert daily["Sign-in recency"] == kw.RECENCY_TODAY
    assert daily["Current streak (days)"] == 20 and daily["Longest streak (days)"] == 20
    assert daily["Mock interview · last 30 days"] == 2 and daily["Mock interview · all time"] == 4
    assert daily["Mock interviews completed · all time"] == 3
    assert daily["Features used (last 30 days)"] == 3 and daily["Used anything (last 30 days)"] == kw.YES
    assert daily["LinkedIn status"] == kw.LINKEDIN_PERSONAL
    assert daily["Completeness band"] == kw.COMPLETE_BARELY  # 3 of 12
    assert occasional["Sign-in days (last 30 days)"] == 2
    assert occasional["Sign-in frequency"] == kw.FREQ_1_2
    assert occasional["Sign-in recency"] == kw.RECENCY_WEEK
    assert occasional["Days since last sign-in"] == 3
    assert occasional["Current streak (days)"] == 0 and occasional["Longest streak (days)"] == 2
    assert occasional["LinkedIn status"] == kw.LINKEDIN_OTHER_PAGE
    assert occasional["Used anything (last 30 days)"] == kw.NO
    assert never["Sign-in recency"] == kw.RECENCY_NEVER
    assert never["Days since last sign-in"] is None
    assert never["Profile on record"] == kw.NO
    assert never["LinkedIn status"] == kw.LINKEDIN_MISSING
    assert never["Completeness band"] == kw.COMPLETE_NONE
    assert never["Profile completeness"] is None
    assert never["Seated in a batch"] == kw.NO


def test_the_registrations_sheet_decides_hours_and_waiting_the_way_the_queue_does(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    ws = wb[kw.SHEET_DATA_REGISTRATIONS]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    by_status = {r["Status"]: r for r in rows}
    assert set(by_status) == {"Approved by the office", "Rejected", "Waiting for review", "Approved by rule"}
    assert by_status["Approved by the office"]["Hours to decision"] == 24.0
    assert by_status["Approved by the office"]["Days waiting"] is None
    assert by_status["Approved by rule"]["Hours to decision"] == 0.0
    assert by_status["Waiting for review"]["Hours to decision"] is None
    assert by_status["Waiting for review"]["Days waiting"] == 2
    assert by_status["Waiting for review"]["Department"] == kw.NOT_STATED
    assert by_status["Approved by rule"]["LinkedIn given"] == kw.NO
    assert by_status["Rejected"]["Repeat application"] == kw.YES
    assert all(isinstance(r["Month"], datetime) and r["Month"].day == 1 for r in rows)


def test_the_sign_ins_sheet_is_in_the_programmes_clock(workbook_paths) -> None:
    wb = load_workbook(workbook_paths[0])
    ws = wb[kw.SHEET_DATA_SIGN_INS]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 3
    hours = sorted(r["Hour"] for r in rows)
    assert hours == [9, 14, 21]  # the hours they were written at, in IST, not shifted to UTC
    assert {r["Door"] for r in rows} == {"Google", "Password", "Emailed code"}
    assert {r["Device"] for r in rows} == {kw.DEVICE_ANDROID, kw.DEVICE_WINDOWS, kw.DEVICE_OTHER}
    assert all(r["Weekday"] in kw.WEEKDAYS for r in rows)


def test_an_empty_deployment_still_gets_a_workbook(tmp_path) -> None:
    """Zero students is a real state (a college on its first day), and the
    formulas must not divide by it: every share reads blank, not #DIV/0!."""
    path = kw.write_workbook(_snapshot(), tmp_path / "empty.xlsx")
    wb = load_workbook(path)
    assert wb[kw.SHEET_DATA_STUDENTS].max_row == 1
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "/" in cell.value and cell.value.startswith("="):
                    assert "IF(" in cell.value, (ws.title, cell.coordinate, cell.value)


def test_main_refuses_to_overwrite_before_touching_the_database(tmp_path, monkeypatch) -> None:
    existing = tmp_path / "kpis.xlsx"
    existing.write_bytes(b"not a workbook")

    def boom():  # pragma: no cover - the point is that it is never called
        raise AssertionError("the database was opened before the overwrite check")

    monkeypatch.setattr(kw, "SessionLocal", boom)
    assert kw.main(["--out", str(existing)]) == 1
    assert existing.read_bytes() == b"not a workbook"


def test_main_refuses_a_window_it_cannot_mean() -> None:
    with pytest.raises(SystemExit):
        kw.main(["--days", "0", "--out", "x.xlsx"])
    with pytest.raises(SystemExit):
        kw.main(["--days", "400", "--out", "x.xlsx"])


# --------------------------------------------------------------------------- #
# Against the real schema, when there is one.
# --------------------------------------------------------------------------- #


@requires_db
def test_the_loader_reads_the_seeded_database_and_the_writer_takes_it(tmp_path) -> None:
    from app.db import SessionLocal

    db = SessionLocal()
    try:
        snapshot = kw.load_snapshot(db, window_days=30)
    finally:
        db.close()
    assert snapshot.taken_at.tzinfo is not None
    assert snapshot.window_days == 30
    user_ids = {s.user_id for s in snapshot.students}
    assert len(user_ids) == len(snapshot.students)
    assert all(uid in user_ids for uid, _ in snapshot.login_days)
    assert all(e.user_id in user_ids for e in snapshot.sign_ins)
    assert all(e.at.tzinfo is not None for e in snapshot.sign_ins)
    for s in snapshot.students:
        assert set(s.feature_uses) == {f.key for f in kw.FEATURES}
        assert s.stage in kw.STAGE_LABELS
    for r in snapshot.registrations:
        assert r.status in kw.REGISTRATION_STATUS_LABELS
    path = kw.write_workbook(snapshot, tmp_path / "seeded.xlsx", roster=True)
    assert load_workbook(path)[kw.SHEET_DATA_STUDENTS].max_row == len(snapshot.students) + 1
