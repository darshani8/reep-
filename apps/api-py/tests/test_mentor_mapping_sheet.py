"""The mentor list as a spreadsheet (2026-10-10): `app/mentor_mapping.py` and the
three `/admin/mentor-mapping/*` endpoints in `routers/admin_mentoring.py`.

The matcher is pinned without a database, because every rule in it is a
decision somebody will want to change ("surely 'Kiran' should match 'Kiran B'")
and the test is where the decision is written down. The endpoints are pinned
against the real roster for the reason every other mentor-mapping test is: the
write is rule 2's scope key, and what matters is that a file of eighty
assignments goes through exactly the helpers the single assignment goes
through — the group, the college fence, the spell, the audit row — and not
through a sixth writer with rules of its own.
"""

from __future__ import annotations

import io
import uuid

from sqlalchemy import func, select

from conftest import requires_db

from app import mentor_mapping as mm
from app.db import SessionLocal
from app.models.mentor_assignment import MentorAssignment
from app.models.redesign import AuditEvent
from app.models.user import Mentor, Role, Student, User
from app.routers.admin_mentoring import (
    MAPPING_COLUMNS,
    V_ASSIGN,
    V_BLANK,
    V_DUPLICATE,
    V_FACULTY_DISABLED,
    V_FACULTY_NOT_FOUND,
    V_LISTED_TWICE,
    V_STUDENT_AMBIGUOUS,
    V_STUDENT_NOT_FOUND,
    V_UNCHANGED,
)

ADMIN_API = "/api/admin"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER = ("Student Name", "USN", "Mentor and Guide", "Specialization")
REASON = "Mentor and Guide Allocation 2026-27, as approved by the department."


# ------------------------------------------------------------ the matcher --


def _people():
    return [
        mm.Person("dm", "Harshitha N", "1MP25MDM01", "Digital Marketing MBA · 2026-28"),
        mm.Person("hr", "Harshitha N", "1MP25MBA40", "General MBA - Human Resources · 2026-28"),
        mm.Person("ys", "Yashwanth S", "1MP25MBA12", "General MBA - Finance · 2026-28"),
        mm.Person("kb", "Kiran B R", "1MP25MBA13", "General MBA - Finance · 2026-28"),
        mm.Person("kr", "Kiran Rao", "1MP25MBA14", "General MBA - Finance · 2026-28"),
    ]


def test_normalise_strips_titles_and_punctuation_from_the_front_only():
    assert mm.normalise("Prof. Suneel Rao") == "suneel rao"
    assert mm.normalise("Dr. Dr Lalitha Jagadish") == "lalitha jagadish"
    assert mm.normalise("Yashwanth.S") == "yashwanth s"
    assert mm.normalise("  HARSHITHA  N ") == "harshitha n"
    # A surname that is also a title stays: only the FRONT is stripped.
    assert mm.normalise("Anand Sir") == "anand sir"
    assert mm.normalise(None) == ""


def test_a_name_matches_exactly_loosely_ambiguously_or_not_at_all():
    people = _people()
    assert mm.match_name("yashwanth.s", people).kind == mm.MATCH_EXACT
    assert mm.match_name("Yashwanth", people).kind == mm.MATCH_LOOSE
    assert mm.match_name("Kiran", people).kind == mm.MATCH_AMBIGUOUS
    assert mm.match_name("Harshitha N", people).kind == mm.MATCH_AMBIGUOUS
    assert mm.match_name("Nobody Here", people).kind == mm.MATCH_NONE
    assert mm.match_name("", people).kind == mm.MATCH_NONE
    # An initial alone is not a name: "S" is half the roster.
    assert mm.match_name("S", people).kind == mm.MATCH_NONE
    # The handle is exact and case-blind, never loose.
    assert mm.by_key(" 1mp25mba12 ", people).id == "ys"
    assert mm.by_key("1MP25MBA1", people) is None


def test_the_specialization_column_breaks_a_tie_only_when_it_is_decisive():
    people = _people()
    assert mm.match_name("Harshitha N", people, hint="Digital Marketing").person.id == "dm"
    # "HR" shares no word with "Human Resources"; the tie stands and the office
    # is asked for the USN rather than handed a guess.
    assert mm.match_name("Harshitha N", people, hint="HR and Analytics").kind == mm.MATCH_AMBIGUOUS
    # A hint that fits both candidates equally resolves nothing.
    assert mm.match_name("Kiran", people, hint="Core Finance").kind == mm.MATCH_AMBIGUOUS
    # A hint never turns a non-match into a match.
    assert mm.match_name("Nobody", people, hint="Digital Marketing").kind == mm.MATCH_NONE


def test_the_file_resolves_its_own_duplicate_names_and_only_removes_candidates():
    people = _people()
    first = mm.match_name("Harshitha N", people, hint="Digital Marketing")
    second = mm.match_name("Harshitha N", people, hint="HR and Analytics")
    resolved = mm.resolve_claims([first, second])
    assert resolved[0].person.id == "dm"
    assert resolved[1].kind == mm.MATCH_LOOSE and resolved[1].person.id == "hr"
    # Two ambiguous lines with nothing else claiming a candidate stay ambiguous.
    both = mm.resolve_claims([mm.match_name("Kiran", people), mm.match_name("Kiran", people)])
    assert [m.kind for m in both] == [mm.MATCH_AMBIGUOUS, mm.MATCH_AMBIGUOUS]
    assert mm.listed_twice(["a", "b", "a", None, None]) == frozenset({"a"})


# ---------------------------------------------------------- the endpoints --


def _tag() -> str:
    return uuid.uuid4().hex[:6].upper()


def _rename(user_id: str, name: str, usn: str | None = None) -> None:
    with SessionLocal() as db:
        db.get(User, user_id).name = name
        if usn is not None:
            db.scalar(select(Student).where(Student.user_id == user_id)).usn = usn
        db.commit()


def _student_id(user_id: str) -> str:
    with SessionLocal() as db:
        return db.scalar(select(Student.id).where(Student.user_id == user_id))


def _mentor_of(student_id: str) -> str | None:
    with SessionLocal() as db:
        return db.scalar(select(Student.mentor_id).where(Student.id == student_id))


def _group_of(user_id: str) -> str | None:
    with SessionLocal() as db:
        return db.scalar(select(Mentor.id).where(Mentor.user_id == user_id))


def _spells(student_id: str) -> list[MentorAssignment]:
    with SessionLocal() as db:
        return list(
            db.scalars(
                select(MentorAssignment)
                .where(MentorAssignment.student_id == student_id)
                .order_by(MentorAssignment.created_at)
            ).all()
        )


def _audit_rows(student_id: str) -> int:
    with SessionLocal() as db:
        return db.scalar(
            select(func.count()).select_from(AuditEvent).where(
                AuditEvent.entity_id == student_id,
                AuditEvent.action.ilike("%mentor%"),
            )
        )


def _xlsx(rows: list[tuple], header: tuple[str, ...] = HEADER) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _preview(client, headers, payload: bytes, filename: str = "Mentoring_list.xlsx"):
    return client.post(
        f"{ADMIN_API}/mentor-mapping/preview",
        headers=headers,
        files={"file": (filename, payload, XLSX)},
    )


def _apply(client, headers, payload: bytes, reason: str | None = REASON):
    data = {} if reason is None else {"reason": reason}
    return client.post(
        f"{ADMIN_API}/mentor-mapping/apply",
        headers=headers,
        data=data,
        files={"file": ("Mentoring_list.xlsx", payload, XLSX)},
    )


def _by_line(body: dict) -> dict[int, dict]:
    return {row["line_no"]: row for row in body["rows"]}


@requires_db
def test_preview_judges_every_line_and_writes_nothing(client, make_user):
    """One of each answer the judge can give, and the roster untouched after it."""
    tag = _tag()
    admin = make_user("map-adm", Role.ADMIN)
    shakti = make_user("map-f1", Role.MENTOR)
    _rename(shakti.user_id, "Mapping Shakti Kumar")
    dipin = make_user("map-f2", Role.MENTOR)
    _rename(dipin.user_id, "Mapping Dipin R")
    first = make_user("map-s1")
    _rename(first.user_id, "Mapping Harshitha N", usn=f"MAP{tag}A")
    second = make_user("map-s2")
    _rename(second.user_id, "Mapping Harshitha N", usn=f"MAP{tag}B")
    third = make_user("map-s3")
    _rename(third.user_id, "Mapping Yashwanth S", usn=f"MAP{tag}C")
    fourth = make_user("map-s4")
    _rename(fourth.user_id, "Mapping Deepak M N", usn=f"MAP{tag}D")

    payload = _xlsx([
        # 2: the USN settles which Harshitha; the faculty name carries a title.
        ("Mapping Harshitha N", f"MAP{tag}A", "Prof. Mapping Shakti Kumar", "Digital Marketing"),
        # 3: ambiguous on its own, resolved because line 2 claimed the other one;
        #    the faculty member matches loosely ("Mapping Dipin" ⊂ "Mapping Dipin R").
        ("Mapping Harshitha N", None, "Mapping Dipin", "HR and Analytics"),
        # 4: punctuation and case in the name; the faculty member by email.
        ("mapping yashwanth.s", None, dipin.email, None),
        # 5: nobody by that name.
        ("Mapping Nobody At All", None, "Prof. Mapping Shakti Kumar", None),
        # 6: no such faculty account.
        ("Mapping Deepak M N", None, "Prof. Unknown Person", None),
        # 7: a USN nobody holds — the office typed it to be exact, so no name fallback.
        ("Mapping Deepak M N", f"MAP{tag}Z", "Mapping Dipin", None),
        # 8: a line with a faculty member and no student.
        (None, None, "Mapping Dipin", None),
    ])
    r = _preview(client, admin.headers, payload)
    assert r.status_code == 200, r.text
    body = r.json()
    rows = _by_line(body)
    assert rows[2]["verdict"] == V_ASSIGN and rows[2]["student_id"] == _student_id(first.user_id)
    assert rows[2]["faculty_user_id"] == shakti.user_id
    assert rows[3]["verdict"] == V_ASSIGN and rows[3]["student_id"] == _student_id(second.user_id)
    assert rows[3]["faculty_user_id"] == dipin.user_id
    assert rows[4]["verdict"] == V_ASSIGN and rows[4]["student_id"] == _student_id(third.user_id)
    assert rows[4]["faculty_name"] == "Mapping Dipin R"
    assert rows[5]["verdict"] == V_STUDENT_NOT_FOUND
    assert rows[6]["verdict"] == V_FACULTY_NOT_FOUND and "Unknown Person" in rows[6]["detail"]
    assert rows[7]["verdict"] == V_STUDENT_NOT_FOUND and f"MAP{tag}Z" in rows[7]["detail"]
    assert rows[8]["verdict"] == V_BLANK
    assert body["to_assign"] == 3
    assert body["counts"][V_ASSIGN] == 3
    # What the screen prints beside each line: the roster's own name and USN.
    assert rows[4]["student_name"] == "Mapping Yashwanth S"
    assert rows[4]["student_usn"] == f"MAP{tag}C"
    assert all(row["applied"] is False for row in body["rows"])

    # NOTHING WRITTEN: not a pointer, not a spell, not a group.
    for account in (first, second, third, fourth):
        sid = _student_id(account.user_id)
        assert _mentor_of(sid) is None
        assert _spells(sid) == []
    assert _group_of(shakti.user_id) is None and _group_of(dipin.user_id) is None

    # Two same-name students with NO line to tell them apart: ambiguous, with
    # the candidates named so the office can put the USN on the line.
    r = _preview(client, admin.headers, _xlsx([
        ("Mapping Harshitha N", None, "Mapping Dipin", None),
    ]))
    assert r.status_code == 200, r.text
    only = r.json()["rows"][0]
    assert only["verdict"] == V_STUDENT_AMBIGUOUS
    assert {c["detail"].split(" · ")[0] for c in only["candidates"]} == {f"MAP{tag}A", f"MAP{tag}B"}


@requires_db
def test_apply_seats_every_assign_line_through_the_same_writer(client, make_user):
    """The group is created on first use, the spell carries the file's reason,
    the audit row is per student, and a second press of the same file does
    nothing at all — `record_mentor_change` skips an unchanged pair."""
    tag = _tag()
    admin = make_user("map-adm2", Role.ADMIN)
    shakti = make_user("map2-f1", Role.MENTOR)
    _rename(shakti.user_id, "Mapping Shakti Kumar")
    first = make_user("map2-s1")
    _rename(first.user_id, "Mapping Harshitha N", usn=f"MAP{tag}A")
    third = make_user("map2-s3")
    _rename(third.user_id, "Mapping Yashwanth S", usn=f"MAP{tag}C")
    payload = _xlsx([
        ("Mapping Harshitha N", None, "Prof. Mapping Shakti Kumar", "Digital Marketing"),
        ("Mapping Yashwanth S", None, "Prof. Mapping Shakti Kumar", "Core Finance"),
    ])

    r = _apply(client, admin.headers, payload)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["applied"] == 2 and body["reason"] == REASON
    assert all(row["applied"] for row in body["rows"])

    group = _group_of(shakti.user_id)
    assert group is not None, "the first assignment is what makes a faculty member a mentor"
    for account in (first, third):
        sid = _student_id(account.user_id)
        assert _mentor_of(sid) == group
        spells = _spells(sid)
        assert len(spells) == 1 and spells[0].to_at is None
        assert spells[0].mentor_id == group
        assert spells[0].reason == REASON
        assert spells[0].by_user_id == admin.user_id
        assert _audit_rows(sid) == 1

    # The same file again: every line reads "already assigned", nothing moves.
    again = _apply(client, admin.headers, payload)
    assert again.status_code == 200, again.text
    assert again.json()["applied"] == 0
    assert again.json()["counts"] == {V_UNCHANGED: 2}
    for account in (first, third):
        sid = _student_id(account.user_id)
        assert len(_spells(sid)) == 1
        assert _audit_rows(sid) == 1

    # Moving one of them to another faculty member closes the spell and says
    # on the line whom they move from.
    dipin = make_user("map2-f2", Role.MENTOR)
    _rename(dipin.user_id, "Mapping Dipin R")
    moved = _preview(client, admin.headers, _xlsx([
        ("Mapping Yashwanth S", None, "Mapping Dipin R", None),
    ]))
    assert moved.status_code == 200, moved.text
    line = moved.json()["rows"][0]
    assert line["verdict"] == V_ASSIGN
    assert line["current_faculty"] == "Mapping Shakti Kumar"
    assert "Moves from Mapping Shakti Kumar to Mapping Dipin R" in line["detail"]


@requires_db
def test_apply_without_a_reason_is_refused_and_writes_nothing(client, make_user):
    tag = _tag()
    admin = make_user("map-adm3", Role.ADMIN)
    shakti = make_user("map3-f1", Role.MENTOR)
    _rename(shakti.user_id, "Mapping Shakti Kumar")
    first = make_user("map3-s1")
    _rename(first.user_id, "Mapping Harshitha N", usn=f"MAP{tag}A")
    payload = _xlsx([("Mapping Harshitha N", None, "Mapping Shakti Kumar", None)])

    assert _apply(client, admin.headers, payload, reason="   ").status_code == 422
    assert _apply(client, admin.headers, payload, reason=None).status_code == 422
    assert _mentor_of(_student_id(first.user_id)) is None
    assert _group_of(shakti.user_id) is None

    # A file that is not a spreadsheet is refused with the parser's sentence.
    r = client.post(
        f"{ADMIN_API}/mentor-mapping/preview",
        headers=admin.headers,
        files={"file": ("list.pdf", b"%PDF-1.4 not a sheet", "application/pdf")},
    )
    assert r.status_code == 422, r.text
    assert ".xlsx" in r.json()["detail"]


@requires_db
def test_a_faculty_member_by_role_holds_nothing_here(client, make_user):
    """A mentor who could upload a sheet could seat themselves with any student
    in the programme and then read everything about them."""
    mentor = make_user("map-f-only", Role.MENTOR)
    payload = _xlsx([("Mapping Harshitha N", None, "Voice Test map-f-only", None)])
    assert _preview(client, mentor.headers, payload).status_code == 403
    assert _apply(client, mentor.headers, payload).status_code == 403
    r = client.get(f"{ADMIN_API}/mentor-mapping/template.xlsx", headers=mentor.headers)
    assert r.status_code == 403


@requires_db
def test_one_student_on_two_lines_with_two_faculty_members_is_refused(client, make_user):
    """A contradiction the file has to settle; the clean line beside it still lands."""
    tag = _tag()
    admin = make_user("map-adm4", Role.ADMIN)
    shakti = make_user("map4-f1", Role.MENTOR)
    _rename(shakti.user_id, "Mapping Shakti Kumar")
    dipin = make_user("map4-f2", Role.MENTOR)
    _rename(dipin.user_id, "Mapping Dipin R")
    first = make_user("map4-s1")
    _rename(first.user_id, "Mapping Harshitha N", usn=f"MAP{tag}A")
    third = make_user("map4-s3")
    _rename(third.user_id, "Mapping Yashwanth S", usn=f"MAP{tag}C")
    payload = _xlsx([
        ("Mapping Harshitha N", None, "Mapping Shakti Kumar", None),
        ("Mapping Yashwanth S", None, "Mapping Shakti Kumar", None),
        ("Mapping Harshitha N", None, "Mapping Dipin R", None),
        # The same faculty member twice is a repeat, not a contradiction.
        ("Mapping Yashwanth S", None, "Mapping Shakti Kumar", None),
    ])
    r = _apply(client, admin.headers, payload)
    assert r.status_code == 200, r.text
    rows = _by_line(r.json())
    assert rows[2]["verdict"] == V_LISTED_TWICE and rows[4]["verdict"] == V_LISTED_TWICE
    assert "lines 2, 4" in rows[2]["detail"]
    assert rows[3]["verdict"] == V_ASSIGN and rows[3]["applied"]
    assert rows[5]["verdict"] == V_DUPLICATE and not rows[5]["applied"]
    assert r.json()["applied"] == 1
    assert _mentor_of(_student_id(first.user_id)) is None
    assert _mentor_of(_student_id(third.user_id)) == _group_of(shakti.user_id)


@requires_db
def test_a_disabled_faculty_account_or_a_removed_student_is_not_seated(client, make_user):
    tag = _tag()
    admin = make_user("map-adm5", Role.ADMIN)
    shakti = make_user("map5-f1", Role.MENTOR)
    _rename(shakti.user_id, "Mapping Shakti Kumar")
    first = make_user("map5-s1")
    _rename(first.user_id, "Mapping Harshitha N", usn=f"MAP{tag}A")
    gone = make_user("map5-s2")
    _rename(gone.user_id, "Mapping Gone Student", usn=f"MAP{tag}G")
    from datetime import datetime, timezone

    with SessionLocal() as db:
        db.get(User, shakti.user_id).disabled_at = datetime.now(timezone.utc)
        db.get(User, gone.user_id).deleted_at = datetime.now(timezone.utc)
        db.commit()
    r = _preview(client, admin.headers, _xlsx([
        ("Mapping Harshitha N", None, "Mapping Shakti Kumar", None),
        ("Mapping Gone Student", f"MAP{tag}G", "Mapping Shakti Kumar", None),
    ]))
    assert r.status_code == 200, r.text
    rows = _by_line(r.json())
    assert rows[2]["verdict"] == V_FACULTY_DISABLED
    # A removed account is matched by nobody — by USN included.
    assert rows[3]["verdict"] == V_STUDENT_NOT_FOUND


@requires_db
def test_the_template_is_a_file_the_importer_reads(client, make_user):
    """The workbook REEP hands out, filled in and handed back, must judge."""
    import openpyxl

    tag = _tag()
    admin = make_user("map-adm6", Role.ADMIN)
    shakti = make_user("map6-f1", Role.MENTOR)
    _rename(shakti.user_id, "Mapping Shakti Kumar")
    first = make_user("map6-s1")
    _rename(first.user_id, "Mapping Harshitha N", usn=f"MAP{tag}A")

    r = client.get(f"{ADMIN_API}/mentor-mapping/template.xlsx", headers=admin.headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX)
    book = openpyxl.load_workbook(io.BytesIO(r.content))
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    assert header == [name for name, _required in MAPPING_COLUMNS]
    # Replace the example line with a real one and hand it back.
    for column, value in enumerate(
        ("Mapping Harshitha N", f"MAP{tag}A", "Prof. Mapping Shakti Kumar", "Finance"), start=1
    ):
        sheet.cell(row=2, column=column, value=value)
    # The note rows under the example are prose, not students, and read as such.
    buffer = io.BytesIO()
    book.save(buffer)
    judged = _preview(client, admin.headers, buffer.getvalue(), filename="template.xlsx")
    assert judged.status_code == 200, judged.text
    rows = judged.json()["rows"]
    assert rows[0]["verdict"] == V_ASSIGN and rows[0]["student_id"] == _student_id(first.user_id)
    assert judged.json()["to_assign"] == 1
